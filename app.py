import os
import re
import tempfile
import pandas as pd
import pyreadstat
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from itertools import groupby
import streamlit as st

# ==========================================
# PAGE CONFIGURATION
# ==========================================
st.set_page_config(
    page_title="Project Star Cross-Tabulation Generator",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Project Star: Yearly Cross-Tabulation Generator")
st.markdown("Upload your latest yearly SPSS file (`.sav`) below to generate and download the cross-tabulation report.")

# ==========================================
# FILE UPLOADER WIDGET
# ==========================================
uploaded_file = st.file_uploader("Upload Yearly SPSS File (.sav)", type=["sav"])

if uploaded_file is not None:
    if st.button("Generate Yearly Dashboard Report", type="primary"):
        with st.spinner("Processing SPSS file and generating cross-tabulation report... Please wait."):
            
            # Save uploaded file to a temporary file so pyreadstat can read it
            with tempfile.NamedTemporaryFile(delete=False, suffix=".sav") as tmp_file:
                tmp_file.write(uploaded_file.getvalue())
                tmp_file_path = tmp_file.name

            try:
                # ==========================================
                # 1. DATA INGESTION
                # ==========================================
                df, meta = pyreadstat.read_sav(tmp_file_path)
                df = df.loc[:, ~df.columns.duplicated()].copy()

                rm_driver_cols = [f"Q4_{i:02d}" for i in range(1, 24)]
                channel_q2_cols = ['Q2_1', 'Q2_2', 'Q2_3', 'Q2_4', 'Q2_5', 'Q2_6', 'Q2_7']
                pref_channel_col = 'Q2_2_1'
                personal_banker_col = 'Q4A2'
                pb_sat_col = 'Q4B2'
                branch_service_cols = ['Q5_1', 'Q5_2', 'Q5_3', 'Q5_4', 'Q5_21', 'Q5_22', 'Q5_23', 'Q5_24']
                branch_cols = branch_service_cols + ['Q10_2']
                cc_cols = ['Q7_1', 'Q7_2', 'Q7_3', 'Q7_4', 'Q7_5', 'Q7_21', 'Q7_22', 'Q10_3']

                online_cols = ['Q8_1', 'Q8_2', 'Q8_3', 'Q8_4', 'Q8_21', 'Q8_22', 'Q8_23', 'Q8_24', 'Q8_25', 'Q8_26', 'Q8_27']
                app_cols = ['Q9_1', 'Q9_2', 'Q9_3', 'Q9_4', 'Q9_21', 'Q9_22', 'Q9_23', 'Q9_24', 'Q9_25', 'Q9_26']
                chan_sat_cols = ['Q10_1', 'Q10_2', 'Q10_3', 'Q10_4', 'Q10_5']
                product_sat_cols = ['Q11_1_1', 'Q11_1_2', 'Q11_1_3', 'Q11_1_4', 'Q11_1_5']

                q11a_groups = {
                    'Q11A_1': [('Q11a.1. Lending products - Overdraft', 'Q11A_1_1'), ('Q11a.1. Lending products - Loans', 'Q11A_1_2'), ('Q11a.1. Lending products - Other', 'Q11A_1_3')],
                    'Q11A_2': [('Q11a.2. Transactional products - Business account', 'Q11A_2_1'), ('Q11a.2. Transactional products - Debit card', 'Q11A_2_2'), ('Q11a.2. Transactional products - Credit Card', 'Q11A_2_3'), ('Q11a.2. Transactional products - Other', 'Q11A_2_4')],
                    'Q11A_3': [('Q11a.3. Insurance products - Business credit protection plan', 'Q11A_3_1'), ('Q11a.3. Insurance products - Law-on-call business plan', 'Q11A_3_2'), ('Q11a.3. Insurance products - Other', 'Q11A_3_3')],
                    'Q11A_4': [('Q11a.4. Investment products - Savings', 'Q11A_4_1'), ('Q11A_4. Investment products - Notice deposits', 'Q11A_4_2'), ('Q11A_4. Investment products - Other', 'Q11A_4_3')]
                }

                all_q11a_cols = [col for grp in q11a_groups.values() for _, col in grp]
                expectations_cols = ['Q12_1', 'Q12_2', 'Q12_3']
                q16_col = 'Q16'

                consideration_items = [
                    ("Absa", "Q16_1_1_flag"), ("Capitec", "Q16_1_2_flag"), ("Investec", "Q16_1_8_flag"),
                    ("Mercantile", "Q16_1_9_flag"), ("Nedbank", "Q16_1_3_flag"), ("Sasfin", "Q16_1_10_flag"),
                    ("Standard Bank", "Q16_1_4_flag"), ("Some other business banking offering", "Q16_1_5_flag"),
                    ("I would not consider moving from FNB at all", "Q16_1_6_flag")
                ]
                raw_consideration_cols = ['Q16_1_1', 'Q16_1_2', 'Q16_1_3', 'Q16_1_4', 'Q16_1_5', 'Q16_1_6', 'Q16_1_8', 'Q16_1_9', 'Q16_1_10']
                q5a_cols = ['Q5A_1', 'Q5A_3', 'Q5A_4', 'Q5A_5', 'Q5A_6', 'Q5A_7', 'Q5A_8', 'Q5A_9', 'Q5A_10', 'Q5A_11', 'Q5A_12', 'Q5A_13', 'Q5A_14', 'Q5A_2']

                all_rating_cols = list(set(rm_driver_cols + branch_cols + cc_cols + online_cols + app_cols + chan_sat_cols + product_sat_cols + expectations_cols))
                if pb_sat_col and pb_sat_col not in all_rating_cols:
                    all_rating_cols.append(pb_sat_col)

                channel_items = [
                    ("Q2. Business Manager", "Q2_1_flag"), ("Q2. Branch", "Q2_2_flag"), ("Q2. Contact Centre", "Q2_3_flag"),
                    ("Q2. Online Banking", "Q2_4_flag"), ("Q2. FNB Business Banking App", "Q2_5_flag"), ("Q2. Account fulfilment", "Q6_code1_flag"),
                    ("Q2. Product contact centre", "Q6_code2_flag"), ("Q2. Business Desk", "Q6_code3_flag"), ("Q2. Secure chat", "Q2_7_flag"), ("Q2. None of the above", "Q2_6_flag"),
                ]

                q5a_items = [
                    ("Q5a. Prefer face-to-face interaction", "Q5A_1_flag"), ("Q5a. Required to submit documents", "Q5A_3_flag"),
                    ("Q5a. Card collection", "Q5A_4_flag"), ("Q5a. Card queries", "Q5A_5_flag"),
                    ("Q5a. Issue could not be resolved digitally/limited options on digital channels", "Q5A_6_flag"),
                    ("Q5a. Difficulty using digital channels", "Q5A_7_flag"), ("Q5a. Not enough information on digital channels", "Q5A_8_flag"),
                    ("Q5a. Directed to branch", "Q5A_9_flag"), ("Q5a. Needed help/assistance", "Q5A_10_flag"),
                    ("Q5a. Cash deposit/withdrawal", "Q5A_11_flag"), ("Q5a. To open account", "Q5A_12_flag"),
                    ("Q5a. To get bank statement", "Q5A_13_flag"), ("Q5a. Update personal information/details", "Q5A_14_flag"),
                    ("Q5a. Others; specify", "Q5A_2_flag"),
                ]

                q6_items = [
                    ("FICA, outstanding documents relating to your account", "Q6_code_1"),
                    ("A specific product, e.g. such as Instant Solutions", "Q6_code_2"),
                    ("General enquiries (e.g. account, card & cheque-related)", "Q6_code_3"),
                    ("Don't know / not sure", "Q6_code_4")
                ]

                pref_channel_items = [
                    ("Branch", 1), ("Contact Centre", 2), ("Online Banking", 3), ("FNB Business Banking App", 4),
                    ("Business Manager at the branch", 5), ("Business/RM Manager", 6), ("Secure chat Help note", 7), ("None of the above", 8)
                ]

                personal_banker_items = [
                    ("Yes", 1), ("No", 2), ("Do not have a personal FNB Account", 3), ("Refused to answer", 4)
                ]

                rm_labels = {}
                for col in rm_driver_cols:
                    if col in meta.column_names_to_labels and meta.column_names_to_labels[col]:
                        rm_labels[col] = re.sub(r'^(Q4[\._\s\d]*)+', '', meta.column_names_to_labels[col], flags=re.IGNORECASE).strip()
                    else:
                        rm_labels[col] = f"Statement {col}"

                pb_sat_label = "Overall satisfaction with your Personal Banker"
                if pb_sat_col and pb_sat_col in meta.column_names_to_labels and meta.column_names_to_labels[pb_sat_col]:
                    pb_sat_label = re.sub(r'^(Q4B2[\._\s\d]*)+', '', meta.column_names_to_labels[pb_sat_col], flags=re.IGNORECASE).strip()

                explicit_branch_labels = {
                    'Q5_1': 'Offering you personalized service', 'Q5_2': 'The manner in which you are welcomed and directed',
                    'Q5_3': 'Staff understanding your business banking needs', 'Q5_4': 'Waiting time for service',
                    'Q5_21': 'Operating hours of the branch', 'Q5_22': 'Staff communicating with you in a clear and easily understandable way',
                    'Q5_23': 'Staff being willing to help', 'Q5_24': 'Consistently delivering on promises made to you', 'Q10_2': 'OVERALL Branch experience'
                }
                branch_labels = {col: explicit_branch_labels.get(col, col) for col in branch_cols}

                explicit_cc_labels = {
                    'Q7_1': 'Knowledge and competency', 'Q7_2': 'Taking ownership of your query', 'Q7_3': 'Offering you personalized service',
                    'Q7_4': 'Processing requests accurately', 'Q7_5': 'Providing the correct advice relating to your query',
                    'Q7_21': 'The agent providing the correct advice relating to your enquiry or transaction',
                    'Q7_22': 'The agent delivering on promises made', 'Q10_3': 'OVERALL Contact Centre experience'
                }
                cc_labels = {col: explicit_cc_labels.get(col, col) for col in cc_cols}

                online_labels = {col: meta.column_names_to_labels[col] if (col in meta.column_names_to_labels and meta.column_names_to_labels[col]) else col for col in online_cols}
                app_labels = {col: meta.column_names_to_labels[col] if (col in meta.column_names_to_labels and meta.column_names_to_labels[col]) else col for col in app_cols}

                explicit_chan_sat_items = [
                    ("Q10.1. OVERALL - Business Manager experience", "Q10_1"), ("Q10.2. OVERALL - Branch experience", "Q10_2"),
                    ("Q10.3. OVERALL - Contact Centre experience", "Q10_3"), ("Q10.4. OVERALL - Online Banking experience", "Q10_4"),
                    ("Q10.5. OVERALL - FNB Business Banking App experience?", "Q10_5"),
                ]

                explicit_product_sat_items = [
                    ("Q11.1. FNB Business Lending products", "Q11_1_1"), ("Q11.2. FNB Business Transactional products", "Q11_1_2"),
                    ("Q11.3. FNB Business Insurance products", "Q11_1_3"), ("Q11.4. FNB Business Investment products", "Q11_1_4"),
                    ("Q11.5. FNB Business Forex Products", "Q11_1_5")
                ]

                expectations_items = [
                    ("Q12.1. Your overall level of satisfaction with the products you received from FNB Business?", "Q12_1"),
                    ("Q12.2. Your overall level of satisfaction with FNB Business?", "Q12_2"),
                    ("Q12.3. Your overall level of satisfaction with your BM over the last 3 months?", "Q12_3")
                ]

                cols = ['wave', 'REGION', 'SUBREG', 'SEGMENT', 'Type', 'Q14_1', 'Q14_2', 'Q6', q16_col] + all_rating_cols + channel_q2_cols + q5a_cols + [pref_channel_col, personal_banker_col, pb_sat_col] + raw_consideration_cols + all_q11a_cols
                cols_present = [c for c in cols if c in df.columns]

                df_sub = df[cols_present].copy()
                df_sub = df_sub.loc[:, ~df_sub.columns.duplicated()].copy()

                df_sub['Q6_raw'] = df_sub['Q6'].copy() if 'Q6' in df_sub.columns else None

                for col in ['wave', 'REGION', 'SUBREG', 'SEGMENT', 'Type']:
                    if col in meta.variable_value_labels and col in df_sub.columns:
                        df_sub[col] = df_sub[col].map(meta.variable_value_labels[col]).fillna(df_sub[col])

                df_sub['REGION'] = df_sub['REGION'].fillna("Unspecified")
                df_sub['SUBREG'] = df_sub['SUBREG'].fillna("Unspecified")
                df_sub['SEGMENT'] = df_sub['SEGMENT'].fillna("Unspecified")

                segment_relabel_map = {'MEDIUM TOUCH': 'MEDIUM TOUCH (R10-R60M)', 'HIGH TOUCH': 'HIGH TOUCH (R60-R150M)', 'PREMIUM': 'PREMIUM (R150M+)'}
                df_sub['SEGMENT'] = df_sub['SEGMENT'].apply(lambda x: segment_relabel_map.get(str(x).strip(), str(x).strip()))

                def assign_type(segment_val):
                    seg = str(segment_val).strip().upper()
                    if seg in ['R0M-R1M', 'R1M-R5M', 'R5M-R10M', 'MEDIUM TOUCH (R10-R60M)', 'MEDIUM TOUCH']:
                        return 'Growth'
                    elif seg in ['HIGH TOUCH (R60-R150M)', 'HIGH TOUCH', 'PREMIUM (R150M+)', 'PREMIUM']:
                        return 'R10m+'
                    return 'Growth'

                df_sub['Type'] = df_sub['SEGMENT'].apply(assign_type)

                def get_wave_number(val):
                    match = re.search(r'\d+', str(val))
                    return int(match.group()) if match else 999

                df_sub['wave_num'] = df_sub['wave'].apply(get_wave_number)
                sorted_wave_nums = sorted(df_sub['wave_num'].unique())
                regions = sorted([str(x) for x in df_sub['REGION'].unique() if x != "Unspecified"])

                def clean_rating_score(val):
                    try:
                        fval = float(val)
                        if fval == 11 or fval == 11.0: return None
                        return fval if 1 <= fval <= 10 else None
                    except (ValueError, TypeError):
                        return None

                new_cols = {}
                for col in all_rating_cols:
                    if col in df_sub.columns:
                        new_cols[f"{col}_clean"] = df_sub[col].apply(clean_rating_score)

                new_cols['Q14_1_clean'] = df_sub['Q14_1'].apply(lambda x: x if pd.notnull(x) and x in range(0, 11) else None)
                new_cols['Q14_2_clean'] = df_sub['Q14_2'].apply(lambda x: x if pd.notnull(x) and x in range(0, 11) else None)

                def is_q6_match(val, code_num, desc_text):
                    if pd.isnull(val): return False
                    sval = str(val).strip().lower()
                    return sval == str(code_num) or sval == f"{code_num}.0" or desc_text.lower() in sval

                if 'Q6_raw' in df_sub.columns:
                    q10_3_clean = new_cols.get('Q10_3_clean', pd.Series(index=df_sub.index))
                    new_cols['Q6_code1_clean'] = [q10_3_clean[i] if is_q6_match(v, 1, "Business Account Fulfilment") else None for i, v in enumerate(df_sub['Q6_raw'])]
                    new_cols['Q6_code2_clean'] = [q10_3_clean[i] if is_q6_match(v, 2, "Product Contact Centre") else None for i, v in enumerate(df_sub['Q6_raw'])]
                    new_cols['Q6_code3_clean'] = [q10_3_clean[i] if is_q6_match(v, 3, "Business Desk") else None for i, v in enumerate(df_sub['Q6_raw'])]
                    new_cols['Q6_code1_flag'] = [1 if is_q6_match(v, 1, "Business Account Fulfilment") else 0 for v in df_sub['Q6_raw']]
                    new_cols['Q6_code2_flag'] = [1 if is_q6_match(v, 2, "Product Contact Centre") else 0 for v in df_sub['Q6_raw']]
                    new_cols['Q6_code3_flag'] = [1 if is_q6_match(v, 3, "Business Desk") else 0 for v in df_sub['Q6_raw']]
                    for code_val in [1, 2, 3, 4]:
                        new_cols[f"Q6_code_{code_val}"] = df_sub['Q6_raw'].apply(lambda x: 1 if pd.notnull(x) and float(x) == code_val else 0)

                for q11a_key, sub_items in q11a_groups.items():
                    sub_cols = [c for _, c in sub_items if c in df_sub.columns]
                    if sub_cols: df_sub[f"{q11a_key}_Base"] = df_sub[sub_cols].notnull().any(axis=1).astype(int)
                    for _, col in sub_items:
                        if col in df_sub.columns: new_cols[f"{col}_flag"] = df_sub[col].apply(lambda x: 1 if pd.notnull(x) and float(x) == 1 else 0)

                if q16_col in df_sub.columns:
                    df_sub['Q16_mapped'] = df_sub[q16_col].apply(lambda x: 'Yes' if str(x).strip().lower() in ['1', '1.0', 'yes'] else ('No' if str(x).strip().lower() in ['2', '2.0', 'no'] else None))
                    new_cols['Q16_Yes'] = df_sub['Q16_mapped'].apply(lambda x: 1 if x == 'Yes' else 0)
                    new_cols['Q16_No'] = df_sub['Q16_mapped'].apply(lambda x: 1 if x == 'No' else 0)
                    new_cols['Q16_Base'] = df_sub['Q16_mapped'].apply(lambda x: 1 if pd.notnull(x) else 0)
                else:
                    new_cols['Q16_Yes'], new_cols['Q16_No'], new_cols['Q16_Base'] = 0, 0, 0

                for _, flag_col in consideration_items:
                    raw_col = flag_col.replace('_flag', '')
                    new_cols[flag_col] = df_sub[raw_col].apply(lambda x: 1 if pd.notnull(x) and float(x) == 1 else 0) if raw_col in df_sub.columns else 0

                bm_clean, fnb_clean = new_cols['Q14_2_clean'], new_cols['Q14_1_clean']
                new_cols['BM_Det'] = bm_clean.apply(lambda x: 1 if pd.notnull(x) and x <= 6 else 0)
                new_cols['BM_Pas'] = bm_clean.apply(lambda x: 1 if pd.notnull(x) and 7 <= x <= 8 else 0)
                new_cols['BM_Pro'] = bm_clean.apply(lambda x: 1 if pd.notnull(x) and x >= 9 else 0)
                new_cols['BM_Base'] = bm_clean.apply(lambda x: 1 if pd.notnull(x) else 0)

                new_cols['FNB_Det'] = fnb_clean.apply(lambda x: 1 if pd.notnull(x) and x <= 6 else 0)
                new_cols['FNB_Pas'] = fnb_clean.apply(lambda x: 1 if pd.notnull(x) and 7 <= x <= 8 else 0)
                new_cols['FNB_Pro'] = fnb_clean.apply(lambda x: 1 if pd.notnull(x) and x >= 9 else 0)
                new_cols['FNB_Base'] = fnb_clean.apply(lambda x: 1 if pd.notnull(x) else 0)

                for qcol in channel_q2_cols:
                    if qcol in df_sub.columns: new_cols[f"{qcol}_flag"] = df_sub[qcol].apply(lambda x: 1 if pd.notnull(x) and float(x) == 1 else 0)
                for qcol in q5a_cols:
                    if qcol in df_sub.columns: new_cols[f"{qcol}_flag"] = df_sub[qcol].apply(lambda x: 1 if pd.notnull(x) and float(x) == 1 else 0)
                if pref_channel_col in df_sub.columns:
                    for lbl, code in pref_channel_items: new_cols[f"pref_chan_{code}"] = df_sub[pref_channel_col].apply(lambda x: 1 if pd.notnull(x) and float(x) == code else 0)
                if personal_banker_col in df_sub.columns:
                    for lbl, code in personal_banker_items: new_cols[f"pb_code_{code}"] = df_sub[personal_banker_col].apply(lambda x: 1 if pd.notnull(x) and float(x) == code else 0)

                df_sub = pd.concat([df_sub, pd.DataFrame(new_cols, index=df_sub.index)], axis=1)
                df_sub = df_sub.loc[:, ~df_sub.columns.duplicated()].copy()

                q16a_flag_cols = [flag_col for _, flag_col in consideration_items if flag_col in df_sub.columns]
                df_sub['Q16A_Base'] = (df_sub[q16a_flag_cols].sum(axis=1) > 0).astype(int) if q16a_flag_cols else 0
                df_sub['Channel_Base'] = df_sub[[c for c in channel_q2_cols if c in df_sub.columns]].notnull().any(axis=1).astype(int)
                df_sub['Q5A_Base'] = (df_sub[[f"{q}_flag" for q in q5a_cols if f"{q}_flag" in df_sub.columns]].sum(axis=1) > 0).astype(int)
                df_sub['Q6_Base'] = df_sub['Q6_raw'].apply(lambda x: 1 if pd.notnull(x) else 0)
                df_sub['Pref_Channel_Base'] = df_sub[pref_channel_col].apply(lambda x: 1 if pd.notnull(x) else 0) if pref_channel_col in df_sub.columns else 0
                df_sub['PB_Base'] = df_sub[personal_banker_col].apply(lambda x: 1 if pd.notnull(x) else 0) if personal_banker_col in df_sub.columns else 0

                df_xtab = df_sub[df_sub['Type'].isin(['Growth', 'R10m+'])].copy()

                xtab_base_dict = {
                    'Total_n': ('wave_num', 'count'),
                    'FNB_Base': ('FNB_Base', 'sum'), 'FNB_Det': ('FNB_Det', 'sum'), 'FNB_Pro': ('FNB_Pro', 'sum'),
                    'BM_Base': ('BM_Base', 'sum'), 'BM_Det': ('BM_Det', 'sum'), 'BM_Pro': ('BM_Pro', 'sum'),
                    'Channel_Base': ('Channel_Base', 'sum'), 'Q5A_Base': ('Q5A_Base', 'sum'),
                    'Q6_Base': ('Q6_Base', 'sum'), 'Pref_Channel_Base': ('Pref_Channel_Base', 'sum'), 'PB_Base': ('PB_Base', 'sum'),
                }

                for col_name in ['Q16_Base', 'Q16_Yes', 'Q16_No', 'Q16A_Base']:
                    if col_name in df_xtab.columns: xtab_base_dict[col_name] = (col_name, 'sum')
                for _, flag_col in consideration_items:
                    if flag_col in df_xtab.columns: xtab_base_dict[flag_col] = (flag_col, 'sum')
                for q11a_key in q11a_groups.keys():
                    if f"{q11a_key}_Base" in df_xtab.columns: xtab_base_dict[f"{q11a_key}_Base"] = (f"{q11a_key}_Base", 'sum')
                for col in all_rating_cols:
                    clean_name = f"{col}_clean"
                    if clean_name not in df_xtab.columns: df_xtab[clean_name] = None
                    xtab_base_dict[f"{col}_sum"] = (clean_name, 'sum')
                    xtab_base_dict[f"{col}_count"] = (clean_name, 'count')
                for code_num in [1, 2, 3]:
                    clean_col_name = f"Q6_code{code_num}_clean"
                    if clean_col_name not in df_xtab.columns: df_xtab[clean_col_name] = None
                    xtab_base_dict[f"Q6_code{code_num}_sum"] = (clean_col_name, 'sum')
                    xtab_base_dict[f"Q6_code{code_num}_count"] = (clean_col_name, 'count')
                for _, flag_col in channel_items:
                    if flag_col in df_xtab.columns: xtab_base_dict[flag_col] = (flag_col, 'sum')
                for _, flag_col in q5a_items:
                    if flag_col in df_xtab.columns: xtab_base_dict[flag_col] = (flag_col, 'sum')
                for q11a_key, sub_items in q11a_groups.items():
                    for _, col in sub_items:
                        flag_name = f"{col}_flag"
                        if flag_name in df_xtab.columns: xtab_base_dict[flag_name] = (flag_name, 'sum')
                for code_val in [1, 2, 3, 4]:
                    col_name = f"Q6_code_{code_val}"
                    if col_name in df_xtab.columns: xtab_base_dict[col_name] = (col_name, 'sum')
                for _, code in pref_channel_items:
                    flag_name = f"pref_chan_{code}"
                    if flag_name in df_xtab.columns: xtab_base_dict[flag_name] = (flag_name, 'sum')
                for _, code in personal_banker_items:
                    flag_name = f"pb_code_{code}"
                    if flag_name in df_xtab.columns: xtab_base_dict[flag_name] = (flag_name, 'sum')

                xtab_agg_dict = {k: v for k, v in xtab_base_dict.items() if v[0] in df_xtab.columns}
                summary_xtab_wave_type = df_xtab.groupby(['wave_num', 'Type'], as_index=False).agg(**xtab_agg_dict)
                summary_xtab_wave_region = df_xtab.groupby(['wave_num', 'REGION', 'Type'], as_index=False).agg(**xtab_agg_dict)

                wb_xtab = openpyxl.Workbook()
                ws_xtab = wb_xtab.active
                ws_xtab.title = "Cross Tabulation"
                ws_xtab.views.sheetView[0].showGridLines = True

                ws_xtab_wt_data = wb_xtab.create_sheet(title="_CrossTab_WaveType_Data")
                ws_xtab_wt_data.sheet_state = 'hidden'
                ws_xtab_wt_data.append(list(summary_xtab_wave_type.columns))
                for row in summary_xtab_wave_type.itertuples(index=False): ws_xtab_wt_data.append(list(row))

                ws_xtab_wr_data = wb_xtab.create_sheet(title="_CrossTab_WaveReg_Data")
                ws_xtab_wr_data.sheet_state = 'hidden'
                ws_xtab_wr_data.append(list(summary_xtab_wave_region.columns))
                for row in summary_xtab_wave_region.itertuples(index=False): ws_xtab_wr_data.append(list(row))

                FNB_TEAL, FNB_AMBER, LIGHT_TEAL, LIGHT_AMBER, WHITE = "009B9E", "EA8B24", "E5F5F5", "FDF3E7", "FFFFFF"
                font_title_main = Font(name="Calibri", size=14, bold=True, color=WHITE)
                font_title_sub = Font(name="Calibri", size=10, italic=True, color=WHITE)
                font_header = Font(name="Calibri", size=9, bold=True, color=WHITE)
                font_bold = Font(name="Calibri", size=10, bold=True)
                font_regular = Font(name="Calibri", size=10)
                font_sub_subheader = Font(name="Calibri", size=10, italic=True, bold=True, color="105B5C")
                font_filter_lbl = Font(name="Calibri", size=10, bold=True, color=WHITE)
                font_filter_val = Font(name="Calibri", size=11, bold=True, color="105B5C")

                fill_teal_header = PatternFill(start_color=FNB_TEAL, end_color=FNB_TEAL, fill_type="solid")
                fill_amber_header = PatternFill(start_color=FNB_AMBER, end_color=FNB_AMBER, fill_type="solid")
                fill_light_teal = PatternFill(start_color=LIGHT_TEAL, end_color=LIGHT_TEAL, fill_type="solid")
                fill_light_amber = PatternFill(start_color=LIGHT_AMBER, end_color=LIGHT_AMBER, fill_type="solid")

                thin_border_side = Side(border_style="thin", color="B0C4DE")
                thick_border_side = Side(border_style="medium", color=FNB_TEAL)
                border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
                border_box = Border(left=thick_border_side, right=thick_border_side, top=thick_border_side, bottom=thick_border_side)

                align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
                align_left = Alignment(horizontal="left", vertical="center")

                ws_xtab.merge_cells("B2:E2")
                ws_xtab["B2"] = "PROJECT STAR: WAVE-BASED CROSS TABULATION ANALYSIS"
                ws_xtab["B2"].font, ws_xtab["B2"].fill = font_title_main, fill_teal_header

                ws_xtab.merge_cells("B3:E3")
                ws_xtab["B3"] = "Wave as Main Banner, Analysed by Type and Region Sub-Banners"
                ws_xtab["B3"].font, ws_xtab["B3"].fill = font_title_sub, fill_amber_header

                c_flbl = ws_xtab.cell(row=5, column=2, value="Type Filter:")
                c_flbl.font, c_flbl.fill, c_flbl.alignment, c_flbl.border = font_filter_lbl, fill_teal_header, align_center, border_box

                c_fval = ws_xtab.cell(row=5, column=3, value="All")
                c_fval.font, c_fval.fill, c_fval.alignment, c_fval.border = font_filter_val, fill_light_amber, align_center, border_box

                dv_xtab_type = DataValidation(type="list", formula1='"All,Growth,R10m+"', allow_blank=True)
                ws_xtab.add_data_validation(dv_xtab_type)
                dv_xtab_type.add(ws_xtab["C5"])

                type_sub_categories = ['Growth', 'R10m+']
                region_sub_categories = regions

                xtab_columns = []
                for w_num in sorted_wave_nums:
                    for t_val in type_sub_categories: xtab_columns.append(('Type', w_num, t_val))
                    for r_val in region_sub_categories: xtab_columns.append(('Region', w_num, r_val))

                active_xtab_columns = []
                for col_info in xtab_columns:
                    b_type, w_num, cat_val = col_info
                    sub_df = summary_xtab_wave_type[(summary_xtab_wave_type['wave_num'] == w_num) & (summary_xtab_wave_type['Type'] == cat_val)] if b_type == 'Type' else summary_xtab_wave_region[(summary_xtab_wave_region['wave_num'] == w_num) & (summary_xtab_wave_region['REGION'] == cat_val)]
                    if not sub_df.empty and sub_df['Total_n'].sum() > 0:
                        active_xtab_columns.append(col_info)

                def xtab_wt_sumifs(col_name, w_num, type_val):
                    if col_name not in summary_xtab_wave_type.columns: return "0"
                    col_let = get_column_letter(list(summary_xtab_wave_type.columns).index(col_name) + 1)
                    return f'SUMIFS(_CrossTab_WaveType_Data!{col_let}:{col_let}, _CrossTab_WaveType_Data!A:A, {w_num}, _CrossTab_WaveType_Data!B:B, "{type_val}")'

                nps_section = [("BASE & NET PROMOTER SCORES (NPS)", "banner", None), ("Base Count (n)", "count", "Total_n"), ("FNB Net Promoter Score (NPS)", "nps", ("FNB_Pro", "FNB_Det", "FNB_Base")), ("BM Net Promoter Score (NPS)", "nps", ("BM_Pro", "BM_Det", "BM_Base"))]
                channel_section = [("CHANNEL USAGE", "banner", None), ("Base Count (n)", "count", "Channel_Base")] + [(l, "pct_flag", f) for l, f in channel_items]
                pref_channel_section = [("Preferred Channel usage", "banner", None), ("Base Count (n)", "count", "Pref_Channel_Base")] + [(l, "pref_pct_flag", f"pref_chan_{c}") for l, c in pref_channel_items]
                bm_interaction_xtab_section = [("Business Manager Interaction", "banner", None)] + [(rm_labels[c], "mean", (f"{c}_sum", f"{c}_count")) for c in rm_driver_cols]
                personal_banker_section = [("PERSONAL BANKER", "banner", None), ("Base Count (n)", "count", "PB_Base")] + [(l, "pb_pct_flag", f"pb_code_{c}") for l, c in personal_banker_items]
                pb_sat_section = [("PERSONAL BANKER OVERALL SATISFACTION RATING", "banner", None), (pb_sat_label, "mean", (f"{pb_sat_col}_sum", f"{pb_sat_col}_count"))]
                branch_service_xtab_section = [("Branch Service Aspects", "banner", None)] + [(branch_labels[c], "mean", (f"{c}_sum", f"{c}_count")) for c in branch_service_cols]
                q5a_xtab_section = [("Motivations for Choosing In-Branch over Digital Channels/ Call centre", "banner", None), ("Base Count (n)", "count", "Q5A_Base")] + [(l, "q5a_pct_flag", f) for l, f in q5a_items]
                q6_xtab_section = [("Contact Centre Service Aspects", "banner", None), ("Base Count (n)", "count", "Q6_Base")] + [(l, "q6_pct_flag", c) for l, c in q6_items] + [("Contact Centre agent ratings", "sub_header", None)] + [(cc_labels[c], "mean", (f"{c}_sum", f"{c}_count")) for c in cc_cols]
                online_xtab_section = [("Online Banking through laptop or desktop PC Service Aspects", "banner", None)] + [(online_labels[c], "mean", (f"{c}_sum", f"{c}_count")) for c in online_cols]
                app_xtab_section = [("Banking App Service Aspects", "banner", None)] + [(app_labels[c], "mean", (f"{c}_sum", f"{c}_count")) for c in app_cols]
                overall_ratings_xtab_section = [("OVERALL ratings", "banner", None)] + [(l, "mean", (f"{v}_sum", f"{v}_count")) for l, v in explicit_chan_sat_items]
                product_sat_xtab_section = [("Satisfaction with products", "banner", None)] + [(l, "mean", (f"{v}_sum", f"{v}_count")) for l, v in explicit_product_sat_items] + [("Drivers of Dissatisfaction", "sub_header", None)]
                for k, items in q11a_groups.items():
                    product_sat_xtab_section.append((f"Base Count (n) - {k}", "count", f"{k}_Base"))
                    for lbl, code in items: product_sat_xtab_section.append((lbl, "q11a_pct_flag", (code, f"{k}_Base")))
                expectations_xtab_section = [("Satisfaction: Quality of service and product solutions", "banner", None)] + [(l, "mean", (f"{v}_sum", f"{v}_count")) for l, v in expectations_items]
                consideration_xtab_section = [("Business banking consideration", "banner", None), ("Q16. Consideration to switch", "sub_header", None), ("Base Count (n)", "count", "Q16_Base"), ("Yes (n)", "count", "Q16_Yes"), ("Yes (%)", "q16_pct_flag", "Q16_Yes"), ("No (n)", "count", "Q16_No"), ("No (%)", "q16_pct_flag", "Q16_No"), ("Q16a. Banks/Financial service providers considered", "sub_header", None), ("Base Count (n)", "count", "Q16A_Base")] + [(l, "q16a_pct_flag", f) for l, f in consideration_items]

                full_xtab_sections = [
                    nps_section, channel_section, pref_channel_section, bm_interaction_xtab_section,
                    personal_banker_section, pb_sat_section, branch_service_xtab_section, q5a_xtab_section,
                    q6_xtab_section, online_xtab_section, app_xtab_section, overall_ratings_xtab_section,
                    product_sat_xtab_section, expectations_xtab_section, consideration_xtab_section
                ]

                SLATE_HEADER = "2F4F4F"
                FILL_TEAL, FILL_SLATE = PatternFill(start_color=FNB_TEAL, end_color=FNB_TEAL, fill_type="solid"), PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")

                curr_row, total_cols_span = 7, 2 + len(active_xtab_columns)

                for section in full_xtab_sections:
                    ws_xtab.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=total_cols_span)
                    b_cell = ws_xtab.cell(row=curr_row, column=2, value=section[0][0])
                    b_cell.font, b_cell.fill, b_cell.alignment = Font(name="Calibri", size=10, bold=True, italic=True, color=WHITE), fill_amber_header, align_left
                    curr_row += 1

                    # Tier 1
                    ws_xtab.cell(row=curr_row, column=2, value="Metric / Question Driver").font, ws_xtab.cell(row=curr_row, column=2).fill, ws_xtab.cell(row=curr_row, column=2).alignment, ws_xtab.cell(row=curr_row, column=2).border = font_header, fill_teal_header, align_center, border_cell
                    wave_groups = [(w_num, len(list(group))) for w_num, group in groupby(active_xtab_columns, key=lambda x: x[1])]
                    col_ptr = 3
                    for wave_idx, (w_num, span_len) in enumerate(wave_groups):
                        wave_fill = FILL_TEAL if wave_idx % 2 == 0 else FILL_SLATE
                        ws_xtab.merge_cells(start_row=curr_row, start_column=col_ptr, end_row=curr_row, end_column=col_ptr + span_len - 1)
                        w_cell = ws_xtab.cell(row=curr_row, column=col_ptr, value=f"Wave {w_num}")
                        w_cell.font, w_cell.fill, w_cell.alignment = font_header, wave_fill, align_center
                        for c_i in range(col_ptr, col_ptr + span_len): ws_xtab.cell(row=curr_row, column=c_i).border = border_cell
                        col_ptr += span_len
                    curr_row += 1

                    # Tier 2
                    ws_xtab.cell(row=curr_row, column=2, value="Sub-category").font, ws_xtab.cell(row=curr_row, column=2).fill, ws_xtab.cell(row=curr_row, column=2).alignment, ws_xtab.cell(row=curr_row, column=2).border = font_header, fill_teal_header, align_center, border_cell
                    col_ptr = 3
                    for wave_idx, (w_num, span_len) in enumerate(wave_groups):
                        wave_fill = FILL_TEAL if wave_idx % 2 == 0 else FILL_SLATE
                        wave_cols = active_xtab_columns[col_ptr - 3 : col_ptr - 3 + span_len]
                        sub_groups = [(b_type, len(list(group))) for b_type, group in groupby(wave_cols, key=lambda x: x[0])]
                        sub_ptr = col_ptr
                        for b_type, sub_span in sub_groups:
                            ws_xtab.merge_cells(start_row=curr_row, start_column=sub_ptr, end_row=curr_row, end_column=sub_ptr + sub_span - 1)
                            t_cell = ws_xtab.cell(row=curr_row, column=sub_ptr, value=b_type)
                            t_cell.font, t_cell.fill, t_cell.alignment = font_header, wave_fill, align_center
                            for c_i in range(sub_ptr, sub_ptr + sub_span): ws_xtab.cell(row=curr_row, column=c_i).border = border_cell
                            sub_ptr += sub_span
                        col_ptr += span_len
                    curr_row += 1

                    # Tier 3
                    ws_xtab.cell(row=curr_row, column=2, value="Sub-category").font, ws_xtab.cell(row=curr_row, column=2).fill, ws_xtab.cell(row=curr_row, column=2).alignment, ws_xtab.cell(row=curr_row, column=2).border = font_header, fill_teal_header, align_center, border_cell
                    col_ptr = 3
                    for wave_idx, (w_num, span_len) in enumerate(wave_groups):
                        wave_fill = FILL_TEAL if wave_idx % 2 == 0 else FILL_SLATE
                        for i in range(span_len):
                            col_info = active_xtab_columns[(col_ptr - 3) + i]
                            c = ws_xtab.cell(row=curr_row, column=col_ptr + i, value=col_info[2])
                            c.font, c.fill, c.alignment, c.border = font_header, wave_fill, align_center, border_cell
                        col_ptr += span_len
                    curr_row += 1

                    for label, item_type, cols_ref in section[1:]:
                        if item_type == "sub_header":
                            ws_xtab.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=total_cols_span)
                            sh_cell = ws_xtab.cell(row=curr_row, column=2, value=label)
                            sh_cell.font, sh_cell.fill, sh_cell.alignment = font_sub_subheader, fill_light_teal, align_left
                            for c_idx in range(2, total_cols_span + 1): ws_xtab.cell(row=curr_row, column=c_idx).border = border_cell
                            curr_row += 1
                            continue

                        lbl_cell = ws_xtab.cell(row=curr_row, column=2, value=label)
                        lbl_cell.font = font_bold if "Base" in label or "(n)" in label else font_regular
                        lbl_cell.alignment, lbl_cell.border = align_left, border_cell

                        for i, col_info in enumerate(active_xtab_columns):
                            b_type, w_num, cat_val = col_info
                            cell = ws_xtab.cell(row=curr_row, column=3+i)
                            cell.alignment, cell.border = align_center, border_cell

                            if b_type == 'Type':
                                formula_body = ""
                                if item_type == "count":
                                    formula_body = xtab_wt_sumifs(cols_ref, w_num, cat_val)
                                    cell.number_format, cell.font = "#,##0", font_bold
                                elif item_type == "nps":
                                    pro_col, det_col, base_col = cols_ref
                                    f_pro, f_det, f_base = xtab_wt_sumifs(pro_col, w_num, cat_val), xtab_wt_sumifs(det_col, w_num, cat_val), xtab_wt_sumifs(base_col, w_num, cat_val)
                                    formula_body = f"({f_pro} - {f_det}) / {f_base} * 100"
                                    cell.number_format, cell.font = "0.00", font_bold
                                elif item_type == "mean":
                                    sum_col, cnt_col = cols_ref
                                    formula_body = f"IF({xtab_wt_sumifs(cnt_col, w_num, cat_val)}=0, \"\", {xtab_wt_sumifs(sum_col, w_num, cat_val)}/{xtab_wt_sumifs(cnt_col, w_num, cat_val)})"
                                    cell.number_format = "0.00"
                                elif item_type in ["pct_flag", "pref_pct_flag"]:
                                    base_name = "Channel_Base" if item_type == "pct_flag" else "Pref_Channel_Base"
                                    formula_body = f"{xtab_wt_sumifs(cols_ref, w_num, cat_val)} / {xtab_wt_sumifs(base_name, w_num, cat_val)}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q5a_pct_flag":
                                    formula_body = f"{xtab_wt_sumifs(cols_ref, w_num, cat_val)} / {xtab_wt_sumifs('Q5A_Base', w_num, cat_val)}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q6_pct_flag":
                                    formula_body = f"{xtab_wt_sumifs(cols_ref, w_num, cat_val)} / {xtab_wt_sumifs('Q6_Base', w_num, cat_val)}"
                                    cell.number_format = "0.0%"
                                elif item_type == "pb_pct_flag":
                                    formula_body = f"{xtab_wt_sumifs(cols_ref, w_num, cat_val)} / {xtab_wt_sumifs('PB_Base', w_num, cat_val)}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q11a_pct_flag":
                                    flag_col, base_col = cols_ref
                                    formula_body = f"{xtab_wt_sumifs(f'{flag_col}_flag', w_num, cat_val)} / {xtab_wt_sumifs(base_col, w_num, cat_val)}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q16_pct_flag":
                                    formula_body = f"{xtab_wt_sumifs(cols_ref, w_num, cat_val)} / {xtab_wt_sumifs('Q16_Base', w_num, cat_val)}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q16a_pct_flag":
                                    formula_body = f"{xtab_wt_sumifs(cols_ref, w_num, cat_val)} / {xtab_wt_sumifs('Q16A_Base', w_num, cat_val)}"
                                    cell.number_format = "0.0%"
                                
                                cell.value = f'=IF(OR($C$5="All", $C$5="{cat_val}"), IFERROR({formula_body}, 0), "")'
                            else: # Region
                                formula_body = ""
                                if item_type == "count":
                                    formula_body = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    cell.number_format, cell.font = "#,##0", font_bold
                                elif item_type == "nps":
                                    pro_col, det_col, base_col = cols_ref
                                    f_pro = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(pro_col) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(pro_col) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_det = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(det_col) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(det_col) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_base = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(base_col) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(base_col) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"({f_pro} - {f_det}) / {f_base} * 100"
                                    cell.number_format, cell.font = "0.00", font_bold
                                elif item_type == "mean":
                                    sum_col, cnt_col = cols_ref
                                    f_sum = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(sum_col) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(sum_col) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_cnt = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cnt_col) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cnt_col) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"IF({f_cnt}=0, \"\", {f_sum}/{f_cnt})"
                                    cell.number_format = "0.00"
                                elif item_type in ["pct_flag", "pref_pct_flag"]:
                                    base_name = "Channel_Base" if item_type == "pct_flag" else "Pref_Channel_Base"
                                    f_num = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_den = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(base_name) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(base_name) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"{f_num} / {f_den}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q5a_pct_flag":
                                    f_num = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_den = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index('Q5A_Base') + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index('Q5A_Base') + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"{f_num} / {f_den}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q6_pct_flag":
                                    f_num = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_den = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index('Q6_Base') + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index('Q6_Base') + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"{f_num} / {f_den}"
                                    cell.number_format = "0.0%"
                                elif item_type == "pb_pct_flag":
                                    f_num = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_den = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index('PB_Base') + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index('PB_Base') + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"{f_num} / {f_den}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q11a_pct_flag":
                                    flag_col, base_col = cols_ref
                                    f_num = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(f'{flag_col}_flag') + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(f'{flag_col}_flag') + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_den = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(base_col) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(base_col) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"{f_num} / {f_den}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q16_pct_flag":
                                    f_num = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_den = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index('Q16_Base') + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index('Q16_Base') + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"{f_num} / {f_den}"
                                    cell.number_format = "0.0%"
                                elif item_type == "q16a_pct_flag":
                                    f_num = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index(cols_ref) + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    f_den = f"SUMIFS(_CrossTab_WaveReg_Data!{get_column_letter(list(summary_xtab_wave_region.columns).index('Q16A_Base') + 1)}:{get_column_letter(list(summary_xtab_wave_region.columns).index('Q16A_Base') + 1)}, _CrossTab_WaveReg_Data!A:A, {w_num}, _CrossTab_WaveReg_Data!B:B, \"{cat_val}\", _CrossTab_WaveReg_Data!C:C, IF($C$5=\"All\", \"*\", $C$5))"
                                    formula_body = f"{f_num} / {f_den}"
                                    cell.number_format = "0.0%"

                                cell.value = f'=IFERROR({formula_body}, 0)'

                        curr_row += 1
                    curr_row += 1

                ws_xtab.column_dimensions['B'].width = 65
                for c_idx in range(3, 3 + len(active_xtab_columns)):
                    ws_xtab.column_dimensions[get_column_letter(c_idx)].width = 11

                # Save workbook to buffer
                output_buffer = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
                wb_xtab.save(output_buffer.name)
                output_buffer.close()

                st.success("SUCCESS: Yearly Cross-Tabulation Report generated successfully!")

                with open(output_buffer.name, "rb") as file:
                    st.download_button(
                        label="📥 Download Yearly Dashboard (.xlsx)",
                        data=file,
                        file_name="Star_Yearly_CrossTab.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

            except Exception as e:
                st.error(f"An error occurred during processing: {e}")
            finally:
                if os.path.exists(tmp_file_path):
                    os.remove(tmp_file_path)
